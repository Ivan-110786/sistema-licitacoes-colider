from flask import Flask, render_template, request, redirect, url_for, session, send_file, send_from_directory
import sqlite3
import pandas as pd
import os
from datetime import datetime
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import io
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

app = Flask(__name__)
app.secret_key = 'segredo_prefeitura_colider'
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def get_db():
    conn = sqlite3.connect('database.db')
    conn.row_factory = sqlite3.Row
    return conn

def registar_log_global(acao, detalhes):
    try:
        conn = get_db()
        usuario = session.get('username', request.form.get('username', 'anónimo'))
        ip_cliente = request.remote_addr
        data_hora = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        with conn:
            conn.execute('''
                INSERT INTO logs_auditoria (data_hora, usuario, acao, detalhes)
                VALUES (?, ?, ?, ?)
            ''', (data_hora, f"{usuario} ({ip_cliente})", acao, detalhes))
        conn.close()
    except Exception:
        pass

def enviar_email_recuperacao(destinatario, nova_senha_temporaria):
    remetente = "suporte@colider.mt.gov.br"
    senha_email = "sua_senha_do_email"
    servidor_smtp = "mail.colider.mt.gov.br"
    porta = 465
    
    msg = MIMEMultipart()
    msg['From'] = remetente
    msg['To'] = destinatario
    msg['Subject'] = "Recuperação de Senha - Sistema de Licitações Colíder"
    
    corpo = f"Olá,\n\nFoi solicitada a recuperação de senha para a sua conta no Sistema de Licitações.\nA sua nova senha temporária é: {nova_senha_temporaria}\n\nRecomendamos alterar a senha após efetuar o login."
    msg.attach(MIMEText(corpo, 'plain'))
    
    try:
        server = smtplib.SMTP_SSL(servidor_smtp, porta)
        server.login(remetente, senha_email)
        server.sendmail(remetente, destinatario, msg.as_string())
        server.quit()
        return True
    except Exception as e:
        print(f"Erro ao enviar e-mail via SSL: {e}")
        try:
            server = smtplib.SMTP(servidor_smtp, 587)
            server.starttls()
            server.login(remetente, senha_email)
            server.sendmail(remetente, destinatario, msg.as_string())
            server.quit()
            return True
        except Exception as e2:
            print(f"Erro ao enviar e-mail via STARTTLS (587): {e2}")
            return False

def init_db():
    conn = get_db()
    with conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS itens (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                codigo TEXT NOT NULL,
                tipo TEXT NOT NULL,
                codigo_tce TEXT,
                descricao TEXT NOT NULL,
                detalhamento TEXT,
                grupo TEXT,
                subgrupo TEXT,
                unidade_medida TEXT NOT NULL,
                unidade_fornecimento TEXT,
                elemento_despesa TEXT,
                categoria TEXT NOT NULL,
                status TEXT DEFAULT 'Ativo',
                imagem TEXT
            )
        ''')
        try:
            conn.execute("ALTER TABLE itens ADD COLUMN status TEXT DEFAULT 'Ativo'")
        except sqlite3.OperationalError:
            pass
        try:
            conn.execute("ALTER TABLE itens ADD COLUMN imagem TEXT")
        except sqlite3.OperationalError:
            pass

        conn.execute('''
            CREATE TABLE IF NOT EXISTS logs_auditoria (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data_hora TEXT NOT NULL,
                usuario TEXT NOT NULL,
                acao TEXT NOT NULL,
                detalhes TEXT NOT NULL
            )
        ''')

        conn.execute('''
            CREATE TABLE IF NOT EXISTS licitacoes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                numero TEXT NOT NULL,
                modalidade TEXT NOT NULL,
                lei TEXT DEFAULT 'Lei 14.133/2021' NOT NULL,
                objeto TEXT NOT NULL,
                valor_referencia REAL,
                categoria TEXT NOT NULL,
                status TEXT DEFAULT 'Aberta',
                data_cadastro TEXT DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        # Verificação segura de colunas sem violar restrições do SQLite
        cursor_cols = conn.execute("PRAGMA table_info(licitacoes)").fetchall()
        colunas_existentes = [c['name'] for c in cursor_cols]
        if 'data_cadastro' not in colunas_existentes:
            conn.execute("ALTER TABLE licitacoes ADD COLUMN data_cadastro TEXT")
            conn.execute("UPDATE licitacoes SET data_cadastro = datetime('now') WHERE data_cadastro IS NULL")
        if 'lei' not in colunas_existentes:
            conn.execute("ALTER TABLE licitacoes ADD COLUMN lei TEXT DEFAULT 'Lei 14.133/2021'")

        conn.execute('''
            CREATE TABLE IF NOT EXISTS secretarias (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT UNIQUE NOT NULL,
                secretario TEXT NOT NULL
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS quantitativos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                licitacao_id INTEGER,
                secretaria TEXT NOT NULL,
                secretario_epoca TEXT,
                item_id INTEGER,
                quantidade REAL NOT NULL,
                origem_item TEXT DEFAULT 'Catálogo',
                UNIQUE(licitacao_id, secretaria, item_id),
                FOREIGN KEY(licitacao_id) REFERENCES licitacoes(id),
                FOREIGN KEY(item_id) REFERENCES itens(id)
            )
        ''')
        try:
            conn.execute("ALTER TABLE quantitativos ADD COLUMN origem_item TEXT DEFAULT 'Catálogo'")
        except sqlite3.OperationalError:
            pass

        conn.execute('''
            CREATE TABLE IF NOT EXISTS cotacoes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                licitacao_id INTEGER,
                item_id INTEGER,
                fornecedor TEXT NOT NULL,
                preco_unitario REAL NOT NULL,
                comprovante_imagem TEXT,
                data_cotacao TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(licitacao_id) REFERENCES licitacoes(id),
                FOREIGN KEY(item_id) REFERENCES itens(id)
            )
        ''')
        try:
            conn.execute("ALTER TABLE cotacoes ADD COLUMN comprovante_imagem TEXT")
        except sqlite3.OperationalError:
            pass

        conn.execute('''
            CREATE TABLE IF NOT EXISTS usuarios (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                senha TEXT NOT NULL,
                perfil TEXT NOT NULL,
                secretaria TEXT,
                email TEXT
            )
        ''')
        try:
            conn.execute("ALTER TABLE usuarios ADD COLUMN email TEXT")
        except sqlite3.OperationalError:
            pass
        
        cursor = conn.execute('SELECT COUNT(*) FROM secretarias')
        if cursor.fetchone()[0] == 0:
            secs_iniciais = [
                ('Secretaria de Saúde', 'Dr. Responsável Saúde'),
                ('Secretaria de Educação', 'Prof. Responsável Educação'),
                ('Secretaria de Obras', 'Eng. Responsável Obras'),
                ('Secretaria de Administração', 'Gestor de Administração'),
                ('Gabinete do Prefeito', 'Chefe de Gabinete')
            ]
            conn.executemany('INSERT INTO secretarias (nome, secretario) VALUES (?, ?)', secs_iniciais)

        cursor_u = conn.execute('SELECT COUNT(*) FROM usuarios')
        if cursor_u.fetchone()[0] == 0:
            conn.execute("INSERT INTO usuarios (username, senha, perfil, secretaria, email) VALUES ('admin', '123', 'admin', 'Geral', 'admin@colider.mt.gov.br')")
            conn.execute("INSERT INTO usuarios (username, senha, perfil, secretaria, email) VALUES ('julia', '123', 'secretaria', 'Secretaria de Saúde', 'julia@colider.mt.gov.br')")
            conn.execute("INSERT INTO usuarios (username, senha, perfil, secretaria, email) VALUES ('educacao', '123', 'secretaria', 'Secretaria de Educação', 'educacao@colider.mt.gov.br')")
    conn.close()

@app.route('/login', methods=['GET', 'POST'])
def login():
    erro = None
    if request.method == 'POST':
        username = request.form['username']
        senha = request.form['senha']
        conn = get_db()
        user = conn.execute('SELECT * FROM usuarios WHERE username = ? AND senha = ?', (username, senha)).fetchone()
        conn.close()
        if user:
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['perfil'] = user['perfil']
            session['secretaria'] = user['secretaria']
            registar_log_global('LOGIN COM SUCESSO', f'Utilizador {username} efetuou login.')
            return redirect(url_for('painel'))
        else:
            registar_log_global('FALHA DE LOGIN', f'Tentativa de login falhada para: {username}')
            erro = "Utilizador ou senha inválidos."
    return render_template('login.html', erro=erro)

@app.route('/recuperar-senha', methods=['GET', 'POST'])
def recuperar_senha():
    mensagem = None
    erro = None
    if request.method == 'POST':
        username = request.form['username']
        email = request.form['email']
        conn = get_db()
        user = conn.execute('SELECT * FROM usuarios WHERE username = ? AND email = ?', (username, email)).fetchone()
        
        if user:
            senha_temporaria = "mudar123"
            with conn:
                conn.execute('UPDATE usuarios SET senha = ? WHERE id = ?', (senha_temporaria, user['id']))
            
            sucesso_envio = enviar_email_recuperacao(email, senha_temporaria)
            if sucesso_envio:
                registar_log_global('RECUPERAÇÃO DE SENHA', f'E-mail enviado para {username} ({email})')
                mensagem = f"Um e-mail com as instruções e a nova senha temporária foi enviado para: {email}."
            else:
                erro = "Erro ao enviar e-mail. Verifique o servidor SMTP."
        else:
            erro = "Utilizador ou e-mail não encontrados no sistema."
        conn.close()
    return render_template('recuperar_senha.html', mensagem=mensagem, erro=erro)

@app.route('/logout')
def logout():
    registar_log_global('LOGOUT', f'Utilizador {session.get("username", "desconhecido")} saiu do sistema.')
    session.clear()
    return redirect(url_for('login'))

@app.route('/')
def painel():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    conn = get_db()
    licitacoes = conn.execute('SELECT * FROM licitacoes').fetchall()
    itens = conn.execute('SELECT * FROM itens WHERE status = "Ativo"').fetchall()
    secretarias_db = conn.execute('SELECT * FROM secretarias').fetchall()
    
    status_secretarias = []
    for sec in secretarias_db:
        nome_sec = sec['nome']
        total_lancados = conn.execute('''
            SELECT COUNT(DISTINCT item_id) as qtd FROM quantitativos WHERE secretaria = ?
        ''', (nome_sec,)).fetchone()['qtd']
        status_secretarias.append({
            'secretaria': nome_sec, 
            'secretario': sec['secretario'],
            'total': total_lancados
        })

    conn.close()
    return render_template('painel.html', licitacoes=licitacoes, itens=itens, status_secretarias=status_secretarias)

@app.route('/secretarias', methods=['GET', 'POST'])
def gerenciar_secretarias():
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    if request.method == 'POST':
        nome = request.form['nome']
        secretario = request.form['secretario']
        try:
            with conn:
                conn.execute('INSERT INTO secretarias (nome, secretario) VALUES (?, ?)', (nome, secretario))
            registar_log_global('NOVA SECRETARIA', f'Criada a secretaria {nome}')
        except sqlite3.IntegrityError:
            pass
        return redirect(url_for('gerenciar_secretarias'))
        
    secretarias = conn.execute('SELECT * FROM secretarias').fetchall()
    conn.close()
    return render_template('secretarias.html', secretarias=secretarias)

@app.route('/secretarias/editar/<int:id>', methods=['GET', 'POST'])
def editar_secretaria(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    secretaria = conn.execute('SELECT * FROM secretarias WHERE id = ?', (id,)).fetchone()
    
    if request.method == 'POST':
        novo_nome = request.form['nome']
        novo_secretario = request.form['secretario']
        with conn:
            conn.execute('UPDATE secretarias SET nome = ?, secretario = ? WHERE id = ?', (novo_nome, novo_secretario, id))
        registar_log_global('EDIÇÃO DE SECRETARIA', f'Secretaria ID {id} atualizada para {novo_nome}')
        return redirect(url_for('gerenciar_secretarias'))
        
    conn.close()
    return render_template('secretaria_editar.html', secretaria=secretaria)

@app.route('/secretarias/excluir/<int:id>')
def excluir_secretaria(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    sec = conn.execute('SELECT nome FROM secretarias WHERE id = ?', (id,)).fetchone()
    with conn:
        conn.execute('DELETE FROM secretarias WHERE id = ?', (id,))
    if sec:
        registar_log_global('EXCLUSÃO DE SECRETARIA', f'Excluída a secretaria: {sec["nome"]}')
    conn.close()
    return redirect(url_for('gerenciar_secretarias'))

@app.route('/historico-cotacoes')
def historico_cotacoes():
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    filtro_fornecedor = request.args.get('fornecedor', '')
    query = '''
        SELECT c.*, i.codigo, i.descricao, l.numero as licitacao_num 
        FROM cotacoes c
        JOIN itens i ON c.item_id = i.id
        JOIN licitacoes l ON c.licitacao_id = l.id
    '''
    params = []
    if filtro_fornecedor:
        query += ' WHERE c.fornecedor LIKE ?'
        params.append(f'%{filtro_fornecedor}%')
    query += ' ORDER BY c.id DESC'
    
    historico = conn.execute(query, params).fetchall()
    fornecedores = conn.execute('SELECT DISTINCT fornecedor FROM cotacoes').fetchall()
    conn.close()
    return render_template('historico_cotacoes.html', historico=historico, fornecedores=fornecedores, filtro_fornecedor=filtro_fornecedor)

@app.route('/usuarios', methods=['GET', 'POST'])
def gerenciar_usuarios():
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    if request.method == 'POST':
        username = request.form['username']
        senha = request.form['senha']
        perfil = request.form['perfil']
        secretaria = request.form['secretaria'] if perfil == 'secretaria' else 'Geral'
        email = request.form.get('email', '')
        
        try:
            with conn:
                conn.execute('INSERT INTO usuarios (username, senha, perfil, secretaria, email) VALUES (?, ?, ?, ?, ?)',
                             (username, senha, perfil, secretaria, email))
            registar_log_global('NOVO UTILIZADOR', f'Criado o utilizador {username} com e-mail {email}')
        except sqlite3.IntegrityError:
            pass
        return redirect(url_for('gerenciar_usuarios'))
        
    usuarios = conn.execute('SELECT * FROM usuarios').fetchall()
    secretarias = conn.execute('SELECT * FROM secretarias').fetchall()
    conn.close()
    return render_template('usuarios.html', usuarios=usuarios, secretarias=secretarias)

@app.route('/usuarios/redefinir-senha/<int:id>', methods=['POST'])
def redefinir_senha_usuario(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    nova_senha = request.form.get('nova_senha')
    conn = get_db()
    u = conn.execute('SELECT username FROM usuarios WHERE id = ?', (id,)).fetchone()
    
    if nova_senha and u:
        with conn:
            conn.execute('UPDATE usuarios SET senha = ? WHERE id = ?', (nova_senha, id))
        registar_log_global('REDEFINIÇÃO DE SENHA', f'Administrador redefiniu a senha do utilizador: {u["username"]}')
        
    conn.close()
    return redirect(url_for('gerenciar_usuarios'))

@app.route('/usuarios/excluir/<int:id>')
def excluir_usuario(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    u = conn.execute('SELECT username FROM usuarios WHERE id = ?', (id,)).fetchone()
    with conn:
        conn.execute('DELETE FROM usuarios WHERE id = ? AND id != 1', (id,))
    if u:
        registar_log_global('EXCLUSÃO DE UTILIZADOR', f'Excluído o utilizador: {u["username"]}')
    conn.close()
    return redirect(url_for('gerenciar_usuarios'))

@app.route('/itens', methods=['GET', 'POST'])
def gerenciar_itens():
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    if request.method == 'POST':
        if 'arquivo_excel' in request.files and request.files['arquivo_excel'].filename != '':
            file = request.files['arquivo_excel']
            filepath = os.path.join(UPLOAD_FOLDER, file.filename)
            file.save(filepath)
            
            if file.filename.endswith('.csv'):
                df = pd.read_csv(filepath, encoding='utf-8')
            else:
                df = pd.read_excel(filepath)
            
            with conn:
                for _, row in df.iterrows():
                    conn.execute('''
                        INSERT INTO itens (codigo, tipo, codigo_tce, descricao, detalhamento, grupo, subgrupo, unidade_medida, unidade_fornecimento, elemento_despesa, categoria, status, imagem)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Ativo', ?)
                    ''', (
                        str(row.get('codigo', '')),
                        str(row.get('tipo', 'Produto')),
                        str(row.get('codigo_tce', '')),
                        str(row.get('descricao', '')),
                        str(row.get('detalhamento', '')),
                        str(row.get('grupo', '')),
                        str(row.get('subgrupo', '')),
                        str(row.get('unidade_medida', 'Un')),
                        str(row.get('unidade_fornecimento', '')),
                        str(row.get('elemento_despesa', '')),
                        str(row.get('categoria', 'Insumos de Impressão')),
                        str(row.get('imagem', ''))
                    ))
            registar_log_global('IMPORTAÇÃO EM LOTE', f'Importados {len(df)} itens')
            return redirect(url_for('gerenciar_itens'))
        
        codigo = request.form['codigo']
        tipo = request.form['tipo']
        codigo_tce = request.form['codigo_tce']
        descricao = request.form['descricao']
        detalhamento = request.form['detalhamento']
        grupo = request.form['grupo']
        subgrupo = request.form['subgrupo']
        unidade_medida = request.form['unidade_medida']
        unidade_fornecimento = request.form.get('unidade_fornecimento', '')
        elemento_despesa = request.form['elemento_despesa']
        categoria = request.form['categoria']
        
        caminho_imagem = ''
        if 'imagem_arquivo' in request.files and request.files['imagem_arquivo'].filename != '':
            img_file = request.files['imagem_arquivo']
            filename = datetime.now().strftime('%Y%m%d%H%M%S_') + img_file.filename
            img_path = os.path.join(UPLOAD_FOLDER, filename)
            img_file.save(img_path)
            caminho_imagem = filename
        elif request.form.get('imagem_url'):
            caminho_imagem = request.form.get('imagem_url')

        with conn:
            conn.execute('''
                INSERT INTO itens (codigo, tipo, codigo_tce, descricao, detalhamento, grupo, subgrupo, unidade_medida, unidade_fornecimento, elemento_despesa, categoria, status, imagem)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Ativo', ?)
            ''', (codigo, tipo, codigo_tce, descricao, detalhamento, grupo, subgrupo, unidade_medida, unidade_fornecimento, elemento_despesa, categoria, caminho_imagem))
        
        registar_log_global('CADASTRO DE ITEM', f'Cadastrado o item código {codigo}')
        return redirect(url_for('gerenciar_itens'))
        
    termo_busca = request.args.get('busca', '')
    categoria_filtro = request.args.get('categoria', '')
    
    query = 'SELECT * FROM itens WHERE 1=1'
    params = []
    
    if termo_busca:
        query += ' AND (codigo LIKE ? OR descricao LIKE ?)'
        params.extend([f'%{termo_busca}%', f'%{termo_busca}%'])
        
    if categoria_filtro:
        query += ' AND categoria = ?'
        params.append(categoria_filtro)
        
    query += ' ORDER BY id DESC'
    itens = conn.execute(query, params).fetchall()
    
    categorias_disponiveis = conn.execute('SELECT DISTINCT categoria FROM itens ORDER BY categoria').fetchall()
    
    conn.close()
    return render_template('itens.html', itens=itens, termo_busca=termo_busca, categorias_disponiveis=categorias_disponiveis, categoria_filtro=categoria_filtro)

@app.route('/itens/editar/<int:id>', methods=['GET', 'POST'])
def editar_item(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    item = conn.execute('SELECT * FROM itens WHERE id = ?', (id,)).fetchone()
    
    em_uso_aberta = conn.execute('''
        SELECT COUNT(*) as qtd FROM quantitativos q
        JOIN licitacoes l ON q.licitacao_id = l.id
        WHERE q.item_id = ? AND l.status = 'Aberta'
    ''', (id,)).fetchone()['qtd'] > 0
    
    forcado = request.form.get('forcar_atualizacao', '0') == '1'
    
    if em_uso_aberta and not forcado and request.method == 'POST':
        conn.close()
        return render_template('item_editar.html', item=item, aviso_travamento=True)
    
    if request.method == 'POST':
        codigo = request.form['codigo']
        tipo = request.form['tipo']
        codigo_tce = request.form['codigo_tce']
        descricao = request.form['descricao']
        detalhamento = request.form['detalhamento']
        grupo = request.form['grupo']
        subgrupo = request.form['subgrupo']
        unidade_medida = request.form['unidade_medida']
        unidade_fornecimento = request.form.get('unidade_fornecimento', '')
        elemento_despesa = request.form['elemento_despesa']
        categoria = request.form['categoria']
        
        caminho_imagem = item['imagem']
        if 'imagem_arquivo' in request.files and request.files['imagem_arquivo'].filename != '':
            img_file = request.files['imagem_arquivo']
            filename = datetime.now().strftime('%Y%m%d%H%M%S_') + img_file.filename
            img_path = os.path.join(UPLOAD_FOLDER, filename)
            img_file.save(img_path)
            caminho_imagem = filename
        elif request.form.get('imagem_url'):
            caminho_imagem = request.form.get('imagem_url')

        with conn:
            conn.execute('''
                UPDATE itens SET codigo = ?, tipo = ?, codigo_tce = ?, descricao = ?, detalhamento = ?, 
                               grupo = ?, subgrupo = ?, unidade_medida = ?, unidade_fornecimento = ?, 
                               elemento_despesa = ?, categoria = ?, imagem = ?
                WHERE id = ?
            ''', (codigo, tipo, codigo_tce, descricao, detalhamento, grupo, subgrupo, unidade_medida, 
                  unidade_fornecimento, elemento_despesa, categoria, caminho_imagem, id))
        
        registar_log_global('EDIÇÃO GLOBAL DE ITEM', f'Atualizado globalmente o item código {codigo}')
        conn.close()
        return redirect(url_for('gerenciar_itens'))
        
    conn.close()
    return render_template('item_editar.html', item=item, em_uso_aberta=em_uso_aberta)

@app.route('/itens/status/<int:id>', methods=['POST'])
def alterar_status_item(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
    
    novo_status = request.form.get('status')
    conn = get_db()
    item = conn.execute('SELECT codigo, descricao FROM itens WHERE id = ?', (id,)).fetchone()
    with conn:
        conn.execute('UPDATE itens SET status = ? WHERE id = ?', (novo_status, id))
    conn.close()
    
    acao = 'DESATIVAÇÃO DE ITEM' if novo_status == 'Desabilitado' else 'ATIVAÇÃO DE ITEM'
    registar_log_global(acao, f'Item {item["codigo"]} alterado para {novo_status}')
    return redirect(url_for('gerenciar_itens'))

@app.route('/itens/excluir/<int:id>')
def excluir_item(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    item = conn.execute('SELECT codigo, descricao FROM itens WHERE id = ?', (id,)).fetchone()
    with conn:
        conn.execute('DELETE FROM itens WHERE id = ?', (id,))
    conn.close()
    
    if item:
        registar_log_global('EXCLUSÃO DE ITEM', f'Excluído permanentemente o item {item["codigo"]}')
    return redirect(url_for('gerenciar_itens'))

@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_from_directory(UPLOAD_FOLDER, filename)

@app.route('/itens/exportar/<formato>')
def exportar_itens(formato):
    if 'user_id' not in session:
        return redirect(url_for('login'))
        
    conn = get_db()
    itens = conn.execute('SELECT codigo, tipo, codigo_tce, descricao, detalhamento, grupo, subgrupo, unidade_medida, elemento_despesa, categoria, status FROM itens').fetchall()
    conn.close()
    
    df = pd.DataFrame([dict(row) for row in itens])
    
    if formato == 'excel':
        output = io.BytesIO()
        df.to_excel(output, index=False, sheet_name='Catálogo de Itens')
        output.seek(0)
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='catalogo_itens_colider.xlsx')
    elif formato == 'csv':
        output = io.BytesIO()
        df.to_csv(output, index=False, encoding='utf-8-sig', sep=';')
        output.seek(0)
        return send_file(output, mimetype='text/csv', as_attachment=True, download_name='catalogo_itens_colider.csv')
    elif formato == 'word':
        doc = Document()
        doc.add_heading('Prefeitura Municipal de Colíder - MT', level=1)
        doc.add_heading('Catálogo Oficial de Itens e Materiais', level=2)
        table = doc.add_table(rows=1, cols=4)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Código'
        hdr_cells[1].text = 'Descrição'
        hdr_cells[2].text = 'Categoria'
        hdr_cells[3].text = 'Unidade'
        for item in itens:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item['codigo'])
            row_cells[1].text = str(item['descricao'])
            row_cells[2].text = str(item['categoria'])
            row_cells[3].text = str(item['unidade_medida'])
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document', as_attachment=True, download_name='catalogo_itens_colider.docx')
    elif formato == 'pdf':
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()
        elements.append(Paragraph("<b>Prefeitura Municipal de Colíder - MT</b>", styles['Heading1']))
        elements.append(Paragraph("<b>Catálogo de Itens Homologados</b>", styles['Heading2']))
        elements.append(Spacer(1, 12))
        data = [['Código', 'Descrição do Item', 'Categoria', 'Unid.']]
        for item in itens:
            data.append([str(item['codigo']), str(item['descricao'][:50]), str(item['categoria']), str(item['unidade_medida'])])
        t = Table(data, colWidths=[60, 280, 110, 50])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#003366')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
        ]))
        elements.append(t)
        doc.build(elements)
        output.seek(0)
        return send_file(output, mimetype='application/pdf', as_attachment=True, download_name='catalogo_itens_colider.pdf')
        
    return redirect(url_for('gerenciar_itens'))

@app.route('/logs')
def ver_logs():
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    logs = conn.execute('SELECT * FROM logs_auditoria ORDER BY id DESC').fetchall()
    conn.close()
    return render_template('logs.html', logs=logs)

@app.route('/licitacoes', methods=['GET', 'POST'])
def gerenciar_licitacoes():
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    if request.method == 'POST':
        numero = request.form['numero']
        modalidade = request.form['modalidade']
        lei = request.form.get('lei', 'Lei 14.133/2021')
        objeto = request.form['objeto']
        valor_referencia = float(request.form['valor_referencia'] or 0)
        categoria = request.form['categoria']
        data_cadastro = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        with conn:
            conn.execute('''
                INSERT INTO licitacoes (numero, modalidade, lei, objeto, valor_referencia, categoria, data_cadastro) 
                VALUES (?, ?, ?, ?, ?, ?, ?)
            ''', (numero, modalidade, lei, objeto, valor_referencia, categoria, data_cadastro))
        registar_log_global('NOVA LICITAÇÃO', f'Cadastrada a licitação {numero}')
        return redirect(url_for('gerenciar_licitacoes'))
    
    ano_atual = datetime.now().strftime('%Y')
    ano_filtro = request.args.get('ano', ano_atual)
    termo_busca = request.args.get('busca', '')

    query = 'SELECT * FROM licitacoes WHERE 1=1'
    params = []

    if ano_filtro:
        query += " AND (strftime('%Y', data_cadastro) = ? OR data_cadastro LIKE ?)"
        params.extend([ano_filtro, f"{ano_filtro}%"])

    if termo_busca:
        query += " AND (numero LIKE ? OR objeto LIKE ? OR modalidade LIKE ?)"
        params.extend([f"%{termo_busca}%", f"%{termo_busca}%", f"%{termo_busca}%"])

    query += ' ORDER BY id DESC'
    licitacoes = conn.execute(query, params).fetchall()

    try:
        anos_disponiveis = [r['ano'] for r in conn.execute("SELECT DISTINCT strftime('%Y', data_cadastro) as ano FROM licitacoes WHERE ano IS NOT NULL UNION SELECT '2026' UNION SELECT '2027' ORDER BY ano DESC").fetchall()]
    except sqlite3.OperationalError:
        anos_disponiveis = ['2027', '2026']

    categorias_itens = conn.execute('SELECT DISTINCT categoria FROM itens ORDER BY categoria').fetchall()
    conn.close()
    
    return render_template('licitacoes.html', licitacoes=licitacoes, categorias_itens=categorias_itens, ano_filtro=ano_filtro, termo_busca=termo_busca, anos_disponiveis=anos_disponiveis)

@app.route('/licitacao/<int:id>/secretarias', methods=['GET', 'POST'])
def quantitativos_secretaria(id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
        
    conn = get_db()
    licitacao = conn.execute('SELECT * FROM licitacoes WHERE id = ?', (id,)).fetchone()
    itens = conn.execute('SELECT * FROM itens WHERE categoria = ? AND status = "Ativo"', (licitacao['categoria'],)).fetchall()
    
    if session['perfil'] == 'secretaria':
        secretaria_atual = session['secretaria']
    else:
        secretaria_atual = request.args.get('secretaria', '') or request.form.get('secretaria', '')

    sec_info = conn.execute('SELECT secretario FROM secretarias WHERE nome = ?', (secretaria_atual,)).fetchone()
    secretario_atual_nome = sec_info['secretario'] if sec_info else 'Gestor Responsável'

    if request.method == 'POST' and 'salvar_lote' in request.form:
        if secretaria_atual:
            with conn:
                conn.execute('DELETE FROM quantitativos WHERE licitacao_id = ? AND secretaria = ?', (id, secretaria_atual))
                
                for item in itens:
                    qtd_str = request.form.get(f'qtd_{item["id"]}', '0')
                    try:
                        qtd = float(qtd_str)
                    except ValueError:
                        qtd = 0.0
                    
                    if qtd > 0:
                        conn.execute('''
                            INSERT INTO quantitativos (licitacao_id, secretaria, secretario_epoca, item_id, quantidade, origem_item)
                            VALUES (?, ?, ?, ?, ?, 'Catálogo')
                        ''', (id, secretaria_atual, secretario_atual_nome, item['id'], qtd))
                
                novo_codigo = request.form.get('novo_codigo', '').strip()
                novo_descricao = request.form.get('novo_descricao', '').strip()
                novo_unidade = request.form.get('novo_unidade', 'Un').strip()
                novo_qtd_str = request.form.get('novo_quantidade', '0').strip()
                
                if novo_descricao and novo_qtd_str:
                    try:
                        novo_qtd = float(novo_qtd_str)
                    except ValueError:
                        novo_qtd = 0.0
                        
                    if novo_qtd > 0:
                        cursor = conn.execute('''
                            INSERT INTO itens (codigo, tipo, descricao, unidade_medida, categoria, status)
                            VALUES (?, 'Produto', ?, ?, ?, 'Ativo')
                        ''', (novo_codigo or 'EXTR_SEC', f"[ADICIONAL SECRETARIA] {novo_descricao}", novo_unidade, licitacao['categoria']))
                        novo_item_id = cursor.lastrowid
                        
                        conn.execute('''
                            INSERT OR REPLACE INTO quantitativos (licitacao_id, secretaria, secretario_epoca, item_id, quantidade, origem_item)
                            VALUES (?, ?, ?, ?, ?, 'Adicionado pela Secretaria')
                        ''', (id, secretaria_atual, secretario_atual_nome, novo_item_id, novo_qtd))

            registar_log_global('LANÇAMENTO DE DFD / QUANTITATIVOS', f'Secretaria {secretaria_atual} atualizou DFD com itens adicionais')
            return redirect(url_for('quantitativos_secretaria', id=id))

    quantitativos_cadastrados = {}
    origens_cadastradas = {}
    if secretaria_atual:
        rows = conn.execute('''
            SELECT item_id, quantidade, origem_item FROM quantitativos 
            WHERE licitacao_id = ? AND secretaria = ?
        ''', (id, secretaria_atual)).fetchall()
        for r in rows:
            quantitativos_cadastrados[r['item_id']] = r['quantidade']
            origens_cadastradas[r['item_id']] = r['origem_item']

    lancamentos = conn.execute('''
        SELECT q.id, q.secretaria, q.secretario_epoca, i.codigo, i.descricao, i.unidade_medida, q.quantidade, q.origem_item 
        FROM quantitativos q
        JOIN itens i ON q.item_id = i.id
        WHERE q.licitacao_id = ?
        ORDER BY q.secretaria, i.codigo
    ''', (id,)).fetchall()

    conn.close()
    return render_template('secretaria_lancamento.html', 
                           licitacao=licitacao, 
                           itens=itens, 
                           secretaria_atual=secretaria_atual,
                           secretario_atual_nome=secretario_atual_nome,
                           quantitativos_cadastrados=quantitativos_cadastrados,
                           origens_cadastradas=origens_cadastradas,
                           lancamentos=lancamentos)

@app.route('/licitacao/<int:id>/cotacao', methods=['GET', 'POST'])
def cotacao_balizamento(id):
    if 'user_id' not in session or session['perfil'] != 'admin':
        return redirect(url_for('painel'))
        
    conn = get_db()
    licitacao = conn.execute('SELECT * FROM licitacoes WHERE id = ?', (id,)).fetchone()
    
    if request.method == 'POST':
        item_id = request.form['item_id']
        fornecedor = request.form['fornecedor']
        preco = float(request.form['preco'] or 0)
        with conn:
            conn.execute('INSERT INTO cotacoes (licitacao_id, item_id, fornecedor, preco_unitario) VALUES (?, ?, ?, ?)',
                         (id, item_id, fornecedor, preco))
        registar_log_global('REGISTO DE COTAÇÃO', f'Fornecedor {fornecedor} cotado com preço {preco}')
        return redirect(url_for('cotacao_balizamento', id=id))

    itens_licitacao = conn.execute('''
        SELECT DISTINCT i.id, i.codigo, i.descricao, i.unidade_medida, SUM(q.quantidade) as total_qtde
        FROM quantitativos q
        JOIN itens i ON q.item_id = i.id
        WHERE q.licitacao_id = ?
        GROUP BY i.id
    ''', (id,)).fetchall()

    cotacoes = conn.execute('''
        SELECT c.*, i.descricao FROM cotacoes c
        JOIN itens i ON c.item_id = i.id
        WHERE c.licitacao_id = ?
    ''', (id,)).fetchall()

    balizamento = conn.execute('''
        SELECT i.codigo, i.descricao, i.unidade_medida, AVG(c.preco_unitario) as preco_medio, MIN(c.preco_unitario) as menor_preco
        FROM cotacoes c
        JOIN itens i ON c.item_id = i.id
        WHERE c.licitacao_id = ?
        GROUP BY i.id
    ''', (id,)).fetchall()

    conn.close()
    return render_template('cotacao.html', licitacao=licitacao, itens_licitacao=itens_licitacao, cotacoes=cotacoes, balizamento=balizamento)

@app.route('/fornecedor/cotacoes')
def listar_cotacoes_abertas_fornecedor():
    conn = get_db()
    licitacoes = conn.execute('SELECT * FROM licitacoes WHERE status = "Aberta"').fetchall()
    conn.close()
    return render_template('fornecedor_lista.html', licitacoes=licitacoes)

@app.route('/fornecedor/preencher/<int:id>')
def portal_preencher_cotacao(id):
    conn = get_db()
    licitacao = conn.execute('SELECT * FROM licitacoes WHERE id = ?', (id,)).fetchone()
    
    itens_licitacao = conn.execute('''
        SELECT i.id, i.codigo, i.descricao, i.unidade_medida, i.imagem, SUM(q.quantidade) as total_qtde
        FROM quantitativos q
        JOIN itens i ON q.item_id = i.id
        WHERE q.licitacao_id = ?
        GROUP BY i.id
    ''', (id,)).fetchall()
    
    conn.close()
    return render_template('fornecedor_cotacao.html', licitacao=licitacao, itens_licitacao=itens_licitacao)

@app.route('/fornecedor/salvar-cotacao/<int:id>', methods=['POST'])
def salvar_cotacao_fornecedor(id):
    fornecedor = request.form.get('fornecedor')
    conn = get_db()
    
    with conn:
        for key, value in request.form.items():
            if key.startswith('item_id_'):
                item_id = value
                preco_unitario_str = request.form.get(f'preco_unit_{item_id}', '0')
                try:
                    preco_unitario = float(preco_unitario_str)
                except ValueError:
                    preco_unitario = 0.0
                
                comprovante_nome = ''
                file_key = f'comprovante_{item_id}'
                if file_key in request.files and request.files[file_key].filename != '':
                    comp_file = request.files[file_key]
                    comprovante_nome = datetime.now().strftime('%Y%m%d%H%M%S_') + comp_file.filename
                    comp_path = os.path.join(UPLOAD_FOLDER, comprovante_nome)
                    comp_file.save(comp_path)
                
                if preco_unitario > 0:
                    conn.execute('''
                        INSERT INTO cotacoes (licitacao_id, item_id, fornecedor, preco_unitario, comprovante_imagem)
                        VALUES (?, ?, ?, ?, ?)
                    ''', (id, item_id, fornecedor, preco_unitario, comprovante_nome))
                    
    registar_log_global('ENVIO DE COTAÇÃO FORNECEDOR', f'O fornecedor {fornecedor} enviou cotação para a Licitação ID {id}')
    conn.close()
    return render_template('fornecedor_sucesso.html', fornecedor=fornecedor)

if __name__ == '__main__':
    init_db()
    app.run(debug=True, host='0.0.0.0', port=5000)